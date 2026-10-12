import io

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.api.spreadsheet.routes import router as spreadsheet_router


@pytest.fixture
def client() -> TestClient:
    app = FastAPI()
    app.include_router(spreadsheet_router, prefix="/api/v1/spreadsheet")
    return TestClient(app)


def test_generate_styled_excel_endpoint(client: TestClient) -> None:
    payload = {
        "title": "2026年度 Q3 研发能效与运营成本核算表",
        "theme": "financial_emerald",
        "sheet_name": "Q3能效核算",
        "columns": [
            {"key": "team", "label": "研发中心团队", "col_type": "string"},
            {"key": "headcount", "label": "人力编制 (人)", "col_type": "number"},
            {"key": "budget", "label": "算力预算 (万元)", "col_type": "currency"},
            {"key": "completion", "label": "交付达成率", "col_type": "percentage"},
            {"key": "audit_date", "label": "核算完成日期", "col_type": "date"},
        ],
        "rows": [
            {
                "cells": {
                    "team": "智能体执行引擎组",
                    "headcount": 18,
                    "budget": 125.5,
                    "completion": 0.985,
                    "audit_date": "2026-09-30",
                }
            },
            {
                "cells": {
                    "team": "云原生控制平面组",
                    "headcount": 14,
                    "budget": 98.0,
                    "completion": 0.942,
                    "audit_date": "2026-09-30",
                }
            },
        ],
    }

    response = client.post("/api/v1/spreadsheet/generate-excel", json=payload)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "content-disposition" in response.headers
    assert len(response.content) > 1000

    # Parse and verify using openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws.title == "Q3能效核算"

    # Check Title banner row (Row 1)
    assert ws.cell(row=1, column=1).value == "2026年度 Q3 研发能效与运营成本核算表"

    # Check Header row (Row 2)
    assert ws.cell(row=2, column=1).value == "研发中心团队"
    assert ws.cell(row=2, column=2).value == "人力编制 (人)"
    assert ws.cell(row=2, column=3).value == "算力预算 (万元)"
    assert ws.cell(row=2, column=4).value == "交付达成率"
    assert ws.cell(row=2, column=5).value == "核算完成日期"

    # Check Data rows (Row 3 & Row 4)
    assert ws.cell(row=3, column=1).value == "智能体执行引擎组"
    assert ws.cell(row=3, column=2).value == 18
    assert ws.cell(row=3, column=3).value == 125.5
    assert ws.cell(row=3, column=5).value == "2026-09-30"

    # Check freeze pane on header
    assert ws.freeze_panes == "A3"


def test_different_themes_generation(client: TestClient) -> None:
    themes = ["business_navy", "dark_gold", "financial_emerald", "tech_slate"]
    for theme in themes:
        payload = {
            "title": f"Theme Test {theme}",
            "theme": theme,
            "columns": [{"key": "id", "label": "ID", "col_type": "number"}],
            "rows": [{"cells": {"id": 1}}],
        }
        resp = client.post("/api/v1/spreadsheet/generate-excel", json=payload)
        assert resp.status_code == 200
        assert len(resp.content) > 1000


def test_sync_office_endpoint(client: TestClient) -> None:
    platforms = [
        ("feishu_bitable", "feishu.cn/base/"),
        ("wecom_doc", "doc.weixin.qq.com/sheet/"),
        ("dingtalk_table", "alidocs.dingtalk.com/table/"),
    ]

    for plat, expected_domain in platforms:
        payload = {
            "platform": plat,
            "table_name": "核心指标看板",
            "columns": [
                {"key": "kpi", "label": "KPI 指标", "col_type": "string"},
                {"key": "val", "label": "达成数值", "col_type": "number"},
            ],
            "rows": [
                {"cells": {"kpi": "用户活跃度", "val": 99.2}},
            ],
        }

        resp = client.post("/api/v1/spreadsheet/sync-office", json=payload)
        assert resp.status_code == 200
        data = resp.json()
        assert data["success"] is True
        assert data["platform"] == plat
        assert data["record_count"] == 1
        assert expected_domain in data["url"]
        assert data["sync_id"].startswith("sync_")
